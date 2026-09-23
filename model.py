"""
الأنموذج التعليمي: قراءة بنك الأخطاء والتمارين، كشف الأخطاء، توليد البرنامج التدريبي
المخصص لكل لاعب، تصديره إلى Word، والإحصاء للمقارنة القبلية/البعدية بين المجموعات.
"""
import io
import math

import numpy as np
import pandas as pd

PHASES = ["تأسيس", "تطوير", "تثبيت"]
PHYS_TYPES = {"قوة", "بلايومتري", "سرعة", "مرونة"}
SCALED_TYPES = {"قوة", "بلايومتري"}

# أسماء المتغيرات البايوميكانيكية بالعربية
METRIC_LABELS = {
    "approach_speed_mps": ("سرعة الاقتراب النهائية", "م/ث"),
    "com_lowering_cm": ("خفض مركز الثقل قبل الارتقاء", "سم"),
    "knee_at_plant_deg": ("زاوية ركبة الارتقاء لحظة وضع القدم", "°"),
    "knee_min_deg": ("أقل زاوية لركبة الارتقاء أثناء الارتكاز", "°"),
    "knee_at_takeoff_deg": ("زاوية ركبة الارتقاء لحظة ترك الأرض", "°"),
    "contact_time_s": ("زمن الارتقاء", "ث"),
    "trunk_lean_plant_deg": ("ميل الجذع لحظة وضع القدم", "°"),
    "trunk_lean_takeoff_deg": ("ميل الجذع لحظة الارتقاء", "°"),
    "free_knee_lift_pct": ("ارتفاع ركبة الرجل الحرة عن الورك", "%"),
    "free_knee_angle_deg": ("زاوية ركبة الرجل الحرة لحظة الارتقاء", "°"),
    "arms_raised_count": ("عدد الذراعين المرفوعتين لحظة الارتقاء", ""),
    "takeoff_angle_deg": ("زاوية الانطلاق", "°"),
    "hip_rise_cm": ("ارتفاع الورك من الارتقاء إلى القمة", "سم"),
    "time_to_peak_s": ("الزمن من الارتقاء إلى القمة", "ث"),
    "hip_angle_at_peak_deg": ("زاوية الورك عند القمة", "°"),
    "knee_angle_at_peak_deg": ("زاوية الركبتين عند القمة", "°"),
}


# ============================================================================
# 1) بنك الأنموذج
# ============================================================================
def load_bank(path_or_bytes):
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes
    x = pd.read_excel(src, sheet_name=None)
    need = ["الأخطاء", "التمارين", "الاختبارات", "الإعدادات"]
    missing = [s for s in need if s not in x]
    if missing:
        raise ValueError("ملف الأنموذج ينقصه الأوراق: " + "، ".join(missing))
    errors = x["الأخطاء"].dropna(subset=["رمز_الخطأ"]).copy()
    errors["الأولوية"] = pd.to_numeric(errors["الأولوية"], errors="coerce").fillna(9)
    exercises = x["التمارين"].dropna(subset=["رمز_التمرين"]).copy()
    tests = x["الاختبارات"].dropna(subset=["رمز_الاختبار"]).copy()
    settings = {}
    for _, r in x["الإعدادات"].dropna(subset=["المفتاح"]).iterrows():
        try:
            settings[str(r["المفتاح"]).strip()] = float(r["القيمة"])
        except (TypeError, ValueError):
            settings[str(r["المفتاح"]).strip()] = r["القيمة"]
    defaults = dict(weeks=8, units_per_week=3, unit_minutes=90, foundation_weeks=3, development_weeks=3,
                    max_errors=3, tech_per_unit=2, phys_per_unit=2, bmi_mid=23, bmi_high=25,
                    factor_mid=0.9, factor_high=0.8, hurdle_ratio=0.25, box_ratio=0.2)
    for k, v in defaults.items():
        settings.setdefault(k, v)
    return {"errors": errors, "exercises": exercises, "tests": tests, "settings": settings}


# ============================================================================
# 2) كشف الأخطاء من القياسات
# ============================================================================
def _num(v):
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def detect_errors(metrics, gender, bank):
    """يعيد الأخطاء الآلية المتحققة مرتبة حسب الأولوية."""
    found = []
    for _, e in bank["errors"].iterrows():
        if str(e.get("طريقة_الكشف", "")).strip() != "آلي":
            continue
        var = str(e.get("المتغير", "")).strip()
        val = _num(metrics.get(var))
        thr = _num(e.get("الحد_إناث") if gender == "أنثى" else e.get("الحد_ذكور"))
        op = str(e.get("الشرط", "")).strip()
        if val is None or thr is None or op not in ("<", ">", "<=", ">="):
            continue
        hit = {"<": val < thr, ">": val > thr, "<=": val <= thr, ">=": val >= thr}[op]
        if hit:
            found.append({
                "id": e["رمز_الخطأ"], "name": e["الخطأ"], "phase": e.get("المرحلة", ""),
                "variable": var, "value": val, "op": op, "threshold": thr,
                "priority": float(e["الأولوية"]), "description": e.get("الوصف", ""),
                "cause": e.get("السبب_المحتمل", ""),
            })
    return sorted(found, key=lambda d: d["priority"])


def error_info(bank, ids):
    df = bank["errors"]
    rows = df[df["رمز_الخطأ"].isin(ids)].sort_values("الأولوية")
    return rows


# ============================================================================
# 3) البرنامج التدريبي المخصص
# ============================================================================
def personalization(height_cm, weight_kg, s):
    h = float(height_cm) / 100.0
    bmi = float(weight_kg) / (h * h)
    if bmi >= s["bmi_high"]:
        factor = s["factor_high"]
    elif bmi >= s["bmi_mid"]:
        factor = s["factor_mid"]
    else:
        factor = 1.0
    hurdle = int(round(float(height_cm) * s["hurdle_ratio"] / 5.0) * 5)
    box = int(round(float(height_cm) * s["box_ratio"] / 5.0) * 5)
    return {"bmi": round(bmi, 1), "factor": factor, "hurdle_cm": hurdle, "box_cm": box}


def phase_of_week(week, s):
    f, d = int(s["foundation_weeks"]), int(s["development_weeks"])
    if week <= f:
        return "تأسيس"
    if week <= f + d:
        return "تطوير"
    return "تثبيت"


def _dose(ex_row, phase, factor):
    sets = _num(ex_row.get(f"مجموعات_{phase}")) or 1
    reps = _num(ex_row.get(f"تكرار_{phase}")) or 1
    if ex_row.get("النوع") in SCALED_TYPES:
        reps = max(1, round(reps * factor))
    return int(sets), int(reps)


def _fill(text, pers):
    return str(text).replace("{hurdle}", str(pers["hurdle_cm"])).replace("{box}", str(pers["box_cm"]))


def build_program(player, error_ids, bank):
    """
    player: dict فيه code, name, height, weight, gender
    error_ids: رموز الأخطاء المستهدفة (آلية أو بصرية)
    """
    s = bank["settings"]
    ex = bank["exercises"]
    pers = personalization(player["height"], player["weight"], s)

    errs = error_info(bank, list(error_ids)).head(int(s["max_errors"]))
    targets = list(errs["رمز_الخطأ"])

    def pool(err, kinds):
        rows = ex[(ex["رمز_الخطأ"] == err) & (ex["النوع"].isin(kinds))]
        return [r for _, r in rows.iterrows()]

    tech = {e: pool(e, {"تكنيك"}) for e in targets}
    phys = {e: pool(e, PHYS_TYPES) for e in targets}
    gen_phys = pool("GEN", PHYS_TYPES)
    warm = pool("GEN", {"إحماء"})
    cool = pool("GEN", {"تهدئة"})
    full = [r for r in pool("GEN", {"تكنيك"})]

    weeks, upw = int(s["weeks"]), int(s["units_per_week"])
    t_per, p_per = int(s["tech_per_unit"]), int(s["phys_per_unit"])

    def item(section, r, phase):
        sets, reps = _dose(r, phase, pers["factor"])
        return {
            "القسم": section, "رمز": r["رمز_التمرين"], "الخطأ": r["رمز_الخطأ"],
            "التمرين": _fill(r["التمرين"], pers), "النوع": r["النوع"],
            "الجرعة": (f"{reps} {r.get('وحدة_التكرار', '')}" if sets == 1 else
                       f"{sets} × {reps} {r.get('وحدة_التكرار', '')}").strip(),
            "الراحة": f"{int(_num(r.get('الراحة_ث')) or 0)} ث" if _num(r.get("الراحة_ث")) else "—",
            "النقطة التعليمية": _fill(r.get("النقطة_التعليمية", "") if pd.notna(r.get("النقطة_التعليمية")) else "", pers),
        }

    def pick(pools_by_err, count, u, fallback):
        chosen, used = [], set()
        order = [e for e in targets if pools_by_err.get(e)]
        k = 0
        while order and len(chosen) < count and k < count * 4:
            e = order[(u + k) % len(order)]
            lst = pools_by_err[e]
            r = lst[(u // max(1, len(order)) + k) % len(lst)]
            if r["رمز_التمرين"] not in used:
                chosen.append(r); used.add(r["رمز_التمرين"])
            k += 1
        i = 0
        while len(chosen) < count and fallback and i < len(fallback) * 2:
            r = fallback[(u + i) % len(fallback)]
            if r["رمز_التمرين"] not in used:
                chosen.append(r); used.add(r["رمز_التمرين"])
            i += 1
        return chosen

    plan = []
    u = 0
    for w in range(1, weeks + 1):
        phase = phase_of_week(w, s)
        units = []
        for j in range(1, upw + 1):
            items = [item("الإحماء", r, phase) for r in warm[:1]]
            items += [item("الجزء الرئيسي: تكنيك تصحيحي", r, phase) for r in pick(tech, t_per, u, [])]
            if phase != "تأسيس" and j == upw and full:
                items.append(item("الجزء الرئيسي: الأداء الكامل", full[0], phase))
            items += [item("الجزء الرئيسي: إعداد بدني خاص", r, phase) for r in pick(phys, p_per, u, gen_phys)]
            items += [item("الختام", r, phase) for r in cool[:1]]
            units.append({"unit": j, "items": items})
            u += 1
        plan.append({"week": w, "phase": phase, "units": units})

    return {
        "player": player,
        "personalization": pers,
        "targets": errs[["رمز_الخطأ", "الخطأ", "المرحلة", "الوصف", "الأولوية"]].to_dict("records"),
        "weeks": plan,
        "settings": {k: s[k] for k in ("weeks", "units_per_week", "unit_minutes")},
    }


def program_table(program):
    rows = []
    for w in program["weeks"]:
        for u in w["units"]:
            for it in u["items"]:
                rows.append({"الأسبوع": w["week"], "المرحلة": w["phase"], "الوحدة": u["unit"], **it})
    return pd.DataFrame(rows)


# ============================================================================
# 4) تصدير البرنامج إلى Word (من اليمين لليسار)
# ============================================================================
def _rtl_paragraph(p):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    for run in p.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rPr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), "Arial")


def _para(doc_or_cell, text, bold=False, size=None, align="right", style=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    run = p.add_run(str(text))
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    if bold:
        rPr.append(OxmlElement("w:bCs"))
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rPr.append(szcs)
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    _rtl_paragraph(p)
    return p


def _table_rtl(table):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tblPr = table._tbl.tblPr
    b = OxmlElement("w:bidiVisual")
    b.set(qn("w:val"), "1")
    tblPr.append(b)


def program_to_docx(program, researcher=""):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.8)
        sectPr = sec._sectPr
        b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); sectPr.append(b)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)

    pl, pers = program["player"], program["personalization"]
    _para(doc, "البرنامج التدريبي التصحيحي — الوثب العالي (فوسبري فلوب)", bold=True, size=16, align="center")
    _para(doc, f"اللاعب: {pl.get('name', '')}  |  الرمز: {pl.get('code', '')}  |  المجموعة: {pl.get('group', '')}")
    _para(doc, f"الطول: {pl['height']} سم  |  الوزن: {pl['weight']} كغم  |  مؤشر كتلة الجسم: {pers['bmi']}  |  "
               f"معامل حجم التمارين: {pers['factor']}")
    st_ = program["settings"]
    _para(doc, f"مدة البرنامج: {int(st_['weeks'])} أسابيع × {int(st_['units_per_week'])} وحدات أسبوعياً "
               f"(زمن الوحدة {int(st_['unit_minutes'])} دقيقة)  |  ارتفاع الحواجز ≈ {pers['hurdle_cm']} سم  |  "
               f"ارتفاع الصندوق ≈ {pers['box_cm']} سم")
    if researcher:
        _para(doc, f"إعداد: {researcher}")

    _para(doc, "الأخطاء الفنية المستهدفة (حسب الأولوية)", bold=True, size=13)
    for t in program["targets"]:
        _para(doc, f"• {t['الخطأ']} ({t['رمز_الخطأ']} — مرحلة {t['المرحلة']}): {t['الوصف']}")
    if not program["targets"]:
        _para(doc, "لم تُحدد أخطاء؛ البرنامج مبني على التمارين العامة.")

    cols = ["القسم", "التمرين", "الجرعة", "الراحة", "النقطة التعليمية"]
    widths = [Cm(3.2), Cm(6.0), Cm(2.6), Cm(1.6), Cm(4.8)]
    for w in program["weeks"]:
        _para(doc, f"الأسبوع {w['week']} — مرحلة {w['phase']}", bold=True, size=13)
        for u in w["units"]:
            _para(doc, f"الوحدة التدريبية {u['unit']}", bold=True)
            tbl = doc.add_table(rows=1, cols=len(cols))
            tbl.style = "Table Grid"
            _table_rtl(tbl)
            for i, c in enumerate(cols):
                cell = tbl.rows[0].cells[i]
                cell.paragraphs[0].text = ""
                cell.width = widths[i]
                r = cell.paragraphs[0].add_run(c); r.bold = True
                r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
                _rtl_paragraph(cell.paragraphs[0])
            for it in u["items"]:
                row = tbl.add_row().cells
                for i, c in enumerate(cols):
                    row[i].width = widths[i]
                    row[i].paragraphs[0].add_run(str(it.get(c, "")))
                    _rtl_paragraph(row[i].paragraphs[0])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============================================================================
# 5) تجميع البيانات والإحصاء
# ============================================================================
def long_data(players, analyses, tests, bank, how="mean"):
    """جدول طويل: الرمز، المجموعة، القياس (قبلي/بعدي)، المتغير، القيمة."""
    frames = []
    groups = players.set_index("code")["group"].to_dict() if not players.empty and "code" in players else {}
    if not analyses.empty:
        keep = [k for k in METRIC_LABELS if k in analyses.columns]
        a = analyses.melt(id_vars=["code", "phase"], value_vars=keep, var_name="var", value_name="value")
        a["value"] = pd.to_numeric(a["value"], errors="coerce")
        a["label"] = a["var"].map(lambda k: METRIC_LABELS[k][0])
        a["unit"] = a["var"].map(lambda k: METRIC_LABELS[k][1])
        a["kind"] = "بايوميكانيكي"
        a["better"] = ""
        frames.append(a)
    if not tests.empty:
        tmeta = bank["tests"].set_index("رمز_الاختبار")
        t = tests[["code", "phase", "test_id", "value"]].rename(columns={"test_id": "var"}).copy()
        t["value"] = pd.to_numeric(t["value"], errors="coerce")
        t["label"] = t["var"].map(lambda k: tmeta["الاختبار"].get(k, k))
        t["unit"] = t["var"].map(lambda k: tmeta["الوحدة"].get(k, ""))
        t["kind"] = t["var"].map(lambda k: "الإنجاز" if tmeta["القابلية"].get(k, "") == "الإنجاز" else "قابلية بدنية")
        t["better"] = t["var"].map(lambda k: tmeta["الأفضل"].get(k, ""))
        frames.append(t)
    if not frames:
        return pd.DataFrame(columns=["code", "group", "phase", "var", "label", "unit", "kind", "better", "value"])
    df = pd.concat(frames, ignore_index=True).dropna(subset=["value"])
    df = df[df["phase"].isin(["قبلي", "بعدي"])]  # قياسات المتابعة تُحلل في صفحة اللاعب فقط
    agg = {"mean": "mean", "max": "max", "min": "min", "last": "last"}[how]
    df = (df.groupby(["code", "phase", "var", "label", "unit", "kind", "better"], as_index=False)["value"].agg(agg))
    df["group"] = df["code"].map(groups).fillna("غير محددة")
    return df


def spss_wide(long_df):
    """صف لكل لاعب، وعمود لكل متغير × قياس (جاهز لـ SPSS)."""
    if long_df.empty:
        return pd.DataFrame()
    d = long_df.copy()
    d["col"] = d["var"] + "_" + d["phase"].map({"قبلي": "pre", "بعدي": "post"}).fillna(d["phase"])
    w = d.pivot_table(index=["code", "group"], columns="col", values="value", aggfunc="first")
    order = []
    for v in dict.fromkeys(d["var"]):
        order += [c for c in (f"{v}_pre", f"{v}_post") if c in w.columns]
    order += [c for c in w.columns if c not in order]
    w = w[order].reset_index()
    w.columns.name = None
    return w


def _desc(x):
    x = np.asarray(x, float)
    x = x[~np.isnan(x)]
    if len(x) == 0:
        return np.nan, np.nan, 0
    return x.mean(), (x.std(ddof=1) if len(x) > 1 else np.nan), len(x)


def group_stats(long_df, var):
    """الإحصاء الوصفي، واختبار t للعينات المترابطة (قبلي/بعدي)، وحجم الأثر، ونسبة التحسن."""
    from scipy import stats

    d = long_df[long_df["var"] == var]
    if d.empty:
        return pd.DataFrame(), {}
    better = d["better"].iloc[0]
    rows = []
    for g, gd in d.groupby("group"):
        piv = gd.pivot_table(index="code", columns="phase", values="value", aggfunc="first")
        pre = piv["قبلي"] if "قبلي" in piv else pd.Series(dtype=float)
        post = piv["بعدي"] if "بعدي" in piv else pd.Series(dtype=float)
        m1, s1, n1 = _desc(pre)
        m2, s2, n2 = _desc(post)
        both = piv.dropna(subset=[c for c in ("قبلي", "بعدي") if c in piv]) if {"قبلي", "بعدي"} <= set(piv) else piv.iloc[0:0]
        t = p = dz = np.nan
        if len(both) >= 2:
            diff = both["بعدي"] - both["قبلي"]
            if diff.std(ddof=1) > 0:
                t, p = stats.ttest_rel(both["بعدي"], both["قبلي"])
                dz = diff.mean() / diff.std(ddof=1)
        imp = np.nan
        if m1 and not np.isnan(m1) and not np.isnan(m2) and m1 != 0:
            imp = (m2 - m1) / abs(m1) * 100 * (-1 if better == "أقل" else 1)
        rows.append({
            "المجموعة": g, "ن": int(max(n1, n2)),
            "قبلي: الوسط": m1, "قبلي: الانحراف": s1, "بعدي: الوسط": m2, "بعدي: الانحراف": s2,
            "قيمة t (مترابطة)": t, "الدلالة p": p, "حجم الأثر dz": dz, "نسبة التحسن %": imp,
        })
    table = pd.DataFrame(rows)

    between = {}
    groups = sorted(d["group"].unique())
    if len(groups) == 2:
        for phase in ("قبلي", "بعدي"):
            a = d[(d["group"] == groups[0]) & (d["phase"] == phase)]["value"].dropna()
            b = d[(d["group"] == groups[1]) & (d["phase"] == phase)]["value"].dropna()
            if len(a) >= 2 and len(b) >= 2 and (a.std() > 0 or b.std() > 0):
                t, p = stats.ttest_ind(a, b, equal_var=False)
                sp = math.sqrt(((len(a) - 1) * a.var(ddof=1) + (len(b) - 1) * b.var(ddof=1)) / (len(a) + len(b) - 2))
                dd = (a.mean() - b.mean()) / sp if sp > 0 else np.nan
                between[phase] = {"t": t, "p": p, "d": dd, "groups": groups}
    return table, between


# ============================================================================
# 6) الاستيراد الجماعي من Excel (قالب فارغ + بيانات تجريبية)
# ============================================================================
P_COLS = ["الرمز", "الاسم", "المجموعة", "الجنس", "الطول", "الوزن", "سنة_الميلاد", "رجل_الارتقاء"]
P_KEYS = ["code", "name", "group", "gender", "height", "weight", "birth_year", "takeoff_leg"]
SHEET_P, SHEET_T, SHEET_A = "اللاعبون", "الاختبارات", "التحليل"


def _test_headers(bank):
    return [f"{r['رمز_الاختبار']} | {r['الاختبار']} ({r['الوحدة']})" for _, r in bank["tests"].iterrows()]


def _metric_headers():
    return [f"{k} | {v[0]}" + (f" ({v[1]})" if v[1] else "") for k, v in METRIC_LABELS.items()]


def sample_data(bank, seed=7):
    """بيانات تقريبية واقعية: 6 تجريبية + 6 ضابطة، قبلي وبعدي. للتجربة فقط."""
    rng = np.random.default_rng(seed)
    players, tests, analyses = [], [], []
    tcodes = list(bank["tests"]["رمز_الاختبار"])
    # (الوسط القبلي، الانحراف، تحسن التجريبية، تحسن الضابطة) — الزمن: التحسن بالنقصان
    base = {"T01": (48, 4, 4.0, 1.5), "T02": (46, 4, 4.0, 1.5), "T03": (245, 12, 9, 3),
            "T04": (14.2, 0.8, 0.8, 0.3), "T05": (3.35, 0.12, -0.12, -0.04), "T06": (11.5, 1.0, 0.9, 0.3),
            "T07": (1.45, 0.15, 0.15, 0.05), "T08": (48, 8, -7, -2), "T09": (8.9, 0.4, -0.35, -0.1),
            "A01": (1.72, 0.06, 0.08, 0.03)}
    mbase = {"approach_speed_mps": (6.3, 0.35, 0.3, 0.1), "com_lowering_cm": (3.5, 1.5, 2.0, 0.5),
             "knee_at_plant_deg": (158, 5, 5, 1.5), "knee_min_deg": (138, 5, 5, 1.5),
             "knee_at_takeoff_deg": (168, 4, 3, 1), "contact_time_s": (0.21, 0.02, -0.025, -0.008),
             "trunk_lean_plant_deg": (-6, 5, -5, -1.5), "trunk_lean_takeoff_deg": (8, 4, -4, -1),
             "free_knee_lift_pct": (-4, 3, 4, 1), "free_knee_angle_deg": (95, 10, -8, -2),
             "arms_raised_count": (1, 0.7, 0.5, 0.1), "takeoff_angle_deg": (44, 4, 3.5, 1),
             "hip_rise_cm": (82, 7, 7, 2.5), "time_to_peak_s": (0.41, 0.03, 0.03, 0.01),
             "hip_angle_at_peak_deg": (158, 6, 4, 1), "knee_angle_at_peak_deg": (110, 10, -5, -1)}
    ind = {}
    for g, pref in (("تجريبية", "E"), ("ضابطة", "C")):
        for k in range(1, 7):
            code = f"{pref}{k}"
            h = int(rng.normal(183, 5)); w = int(round(h - 105 + rng.normal(0, 4)))
            players.append([code, f"لاعب تجريبي {code}", g, "ذكر", h, w, int(rng.integers(2004, 2009)),
                            rng.choice(["اليسرى", "اليمنى"])])
            ind[code] = rng.normal(0, 1)  # مستوى اللاعب
    for p in players:
        code, g = p[0], p[2]
        lvl = ind[code]
        for phase in ("قبلي", "بعدي"):
            row = [code, phase]
            for t in tcodes:
                m, sd, ge, gc = base.get(t, (10, 1, 0.5, 0.2))
                v = m + sd * (0.7 * lvl + 0.5 * rng.normal())
                if phase == "بعدي":
                    v += (ge if g == "تجريبية" else gc) + sd * 0.25 * rng.normal()
                row.append(round(v, 2))
            tests.append(row)
            for att in (1, 2):
                vals, mv = [code, phase, att], {}
                for key, (m, sd, ge, gc) in mbase.items():
                    v = m + sd * (0.5 * lvl + 0.6 * rng.normal())
                    if phase == "بعدي":
                        v += (ge if g == "تجريبية" else gc)
                    if key == "arms_raised_count":
                        v = int(np.clip(round(v), 0, 2))
                    mv[key] = round(float(v), 3 if key.endswith("_s") else 1)
                    vals.append(mv[key])
                vals.append(",".join(e["id"] for e in detect_errors(mv, "ذكر", bank)))
                analyses.append(vals)
    return players, tests, analyses


def make_template(bank, sample=False):
    """sample: False = قالب فارغ، True أو "group" = 12 لاعباً، "single" = لاعب واحد بمحاولات متعددة ومتابعة."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if sample == "single":
        players, tests, analyses = sample_single(bank)
    elif sample:
        players, tests, analyses = sample_data(bank)
        tests = [r[:2] + [""] + r[2:] for r in tests]
        analyses = [r[:2] + [""] + r[2:] for r in analyses]
    else:
        players, tests, analyses = [], [], []
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title, header, rows, note):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[1].height = 60
        for i in range(1, len(header) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 16
        ws.freeze_panes = "C2"
        return ws

    ws = sheet(SHEET_P, P_COLS, players, "")
    for col, opts in (("C", '"تجريبية,ضابطة"'), ("D", '"ذكر,أنثى"'), ("H", '"اليسرى,اليمنى"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True); ws.add_data_validation(dv); dv.add(f"{col}2:{col}500")
    ws = sheet(SHEET_T, ["الرمز", "القياس", "الأسبوع"] + _test_headers(bank), tests, "")
    dv = DataValidation(type="list", formula1='"قبلي,متابعة,بعدي"'); ws.add_data_validation(dv); dv.add("B2:B1000")
    ws = sheet(SHEET_A, ["الرمز", "القياس", "الأسبوع", "المحاولة"] + _metric_headers() + ["الأخطاء (رموز مفصولة بفواصل)"],
               analyses, "")
    dv = DataValidation(type="list", formula1='"قبلي,متابعة,بعدي"'); ws.add_data_validation(dv); dv.add("B2:B2000")

    info = wb.create_sheet("اقرأني", 0)
    info.sheet_view.rightToLeft = True
    info.column_dimensions["A"].width = 110
    lines = [
        ("⚠️ بيانات تقريبية مولّدة للتجربة فقط — لا تُستخدم في الرسالة" if sample else "قالب إدخال البيانات"),
        "",
        "• القياس: قبلي أو متابعة أو بعدي. للمتابعة اكتب رقم الأسبوع في عمود «الأسبوع».",
        "• يمكن تكرار الصف نفسه (اللاعب والقياس) لتسجيل عدة محاولات؛ هذا ضروري عند دراسة لاعب واحد.",
        "",
        "• ورقة «اللاعبون»: لاعب في كل صف. الرمز فريد (مثل E1 للتجريبية وC1 للضابطة).",
        "• ورقة «الاختبارات»: صف لكل لاعب في كل قياس (قبلي/بعدي). اترك الخانة فارغة إن لم يُجرَ الاختبار.",
        "• ورقة «التحليل»: اختيارية؛ للمتغيرات البايوميكانيكية المقاسة خارج التطبيق (مثل Kinovea).",
        "  التحليل بالفيديو داخل التطبيق يُحفظ تلقائياً ولا يحتاج هذه الورقة.",
        "• التزم بالوحدة المكتوبة في عنوان كل عمود (مثلاً الإنجاز بالمتر: 1.75 وليس 175).",
        "• لا تغيّر أسماء الأوراق أو عناوين الأعمدة. يمكنك حذف الأعمدة التي لا تحتاجها.",
        "• الاستيراد يضيف البيانات؛ اللاعب الموجود مسبقاً بنفس الرمز لا يُكرر.",
    ]
    for i, l in enumerate(lines, 1):
        info.cell(row=i, column=1, value=l).alignment = Alignment(horizontal="right", wrap_text=True)
    info["A1"].font = Font(bold=True, size=14, color="C00000" if sample else "000000")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _code_of(header):
    return str(header).split("|")[0].strip()


def parse_import(file_bytes, bank, existing_codes=()):
    """يقرأ ملف القالب ويعيد صفوفاً جاهزة للحفظ مع تقرير بالمشكلات."""
    x = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
    out = {"players": [], "tests": [], "analyses": []}
    report = []
    existing = {str(c).strip().upper() for c in existing_codes}
    groups = {}

    if SHEET_P in x:
        df = x[SHEET_P].dropna(how="all")
        for i, r in df.iterrows():
            code = str(r.get("الرمز", "")).strip().upper()
            if not code or code == "NAN":
                continue
            g = str(r.get("المجموعة", "")).strip()
            if g not in ("تجريبية", "ضابطة"):
                report.append(f"اللاعبون، صف {i + 2}: المجموعة «{g}» غير صحيحة (تجريبية/ضابطة).")
                continue
            groups[code] = g
            if code in existing:
                report.append(f"اللاعب {code} موجود مسبقاً؛ لم يُكرر.")
                continue
            row = dict(zip(P_KEYS, [r.get(c) for c in P_COLS]))
            row["code"] = code
            for k in ("height", "weight", "birth_year"):
                row[k] = _num(row.get(k))
            if not row["height"] or not row["weight"]:
                report.append(f"اللاعب {code}: الطول أو الوزن مفقود.")
            out["players"].append({k: ("" if (v is None or (isinstance(v, float) and math.isnan(v))) else v)
                                   for k, v in row.items()})
            existing.add(code)

    valid_tests = set(bank["tests"]["رمز_الاختبار"])
    if SHEET_T in x:
        df = x[SHEET_T].dropna(how="all")
        tcols = [c for c in df.columns if _code_of(c) in valid_tests]
        for i, r in df.iterrows():
            code = str(r.get("الرمز", "")).strip().upper()
            phase = str(r.get("القياس", "")).strip()
            if not code or code == "NAN":
                continue
            if phase not in ("قبلي", "بعدي", "متابعة"):
                report.append(f"الاختبارات، صف {i + 2}: القياس «{phase}» غير صحيح.")
                continue
            week = _num(r.get("الأسبوع"))
            if code not in existing:
                report.append(f"الاختبارات، صف {i + 2}: اللاعب {code} غير مسجل.")
            for c in tcols:
                v = _num(r.get(c))
                if v is not None:
                    out["tests"].append({"code": code, "group": groups.get(code, ""), "phase": phase,
                                         "week": "" if week is None else int(week),
                                         "test_id": _code_of(c), "value": v, "method": "استيراد Excel"})

    if SHEET_A in x:
        df = x[SHEET_A].dropna(how="all")
        mcols = [c for c in df.columns if _code_of(c) in METRIC_LABELS]
        ecol = next((c for c in df.columns if str(c).startswith("الأخطاء")), None)
        for i, r in df.iterrows():
            code = str(r.get("الرمز", "")).strip().upper()
            phase = str(r.get("القياس", "")).strip()
            if not code or code == "NAN":
                continue
            if phase not in ("قبلي", "بعدي", "متابعة"):
                report.append(f"التحليل، صف {i + 2}: القياس «{phase}» غير صحيح.")
                continue
            week = _num(r.get("الأسبوع"))
            row = {"code": code, "group": groups.get(code, ""), "phase": phase,
                   "week": "" if week is None else int(week),
                   "attempt": int(_num(r.get("المحاولة")) or 1), "video": "استيراد Excel"}
            n = 0
            for c in mcols:
                v = _num(r.get(c))
                row[_code_of(c)] = "" if v is None else v
                n += v is not None
            errs = str(r.get(ecol, "") or "") if ecol else ""
            errs = "" if errs.lower() == "nan" else errs.replace("،", ",").replace(" ", "")
            row["errors_final"] = errs
            row["errors_auto"] = ""
            if n:
                out["analyses"].append(row)
    return out, report


# ============================================================================
# 7) تصميم اللاعب الواحد (Single-subject): تحليل على مستوى المحاولات
# ============================================================================
def directions(bank):
    """اتجاه الأفضل لكل متغير: +1 الأعلى أفضل، -1 الأقل أفضل، 0 غير محدد."""
    d = {}
    for _, e in bank["errors"].iterrows():
        var, op = str(e.get("المتغير", "")).strip(), str(e.get("الشرط", "")).strip()
        if var and op in ("<", "<="):
            d[var] = 1
        elif var and op in (">", ">="):
            d[var] = -1
    for _, t in bank["tests"].iterrows():
        d[t["رمز_الاختبار"]] = 1 if t["الأفضل"] == "أعلى" else -1 if t["الأفضل"] == "أقل" else 0
    return d


def phase_order(phase, week, total_weeks):
    if phase == "قبلي":
        return 0
    if phase == "بعدي":
        return total_weeks + 1
    w = _num(week)
    return w if w is not None else total_weeks / 2


def trials_data(analyses, tests, bank, code):
    """كل محاولة كما هي (بدون تجميع) للاعب واحد."""
    frames = []
    total = int(bank["settings"]["weeks"])
    if not analyses.empty and "code" in analyses:
        a = analyses[analyses["code"].astype(str) == str(code)].copy()
        if not a.empty:
            if "week" not in a:
                a["week"] = ""
            a["trial"] = np.arange(len(a))
            keep = [k for k in METRIC_LABELS if k in a.columns]
            m = a.melt(id_vars=["phase", "week", "trial"], value_vars=keep, var_name="var", value_name="value")
            m["label"] = m["var"].map(lambda k: METRIC_LABELS[k][0])
            m["unit"] = m["var"].map(lambda k: METRIC_LABELS[k][1])
            m["kind"] = "بايوميكانيكي"
            frames.append(m)
    if not tests.empty and "code" in tests:
        t = tests[tests["code"].astype(str) == str(code)].copy()
        if not t.empty:
            if "week" not in t:
                t["week"] = ""
            tm = bank["tests"].set_index("رمز_الاختبار")
            t = t.rename(columns={"test_id": "var"})
            t["trial"] = np.arange(len(t))
            t["label"] = t["var"].map(lambda k: tm["الاختبار"].get(k, k))
            t["unit"] = t["var"].map(lambda k: tm["الوحدة"].get(k, ""))
            t["kind"] = t["var"].map(lambda k: "الإنجاز" if tm["القابلية"].get(k, "") == "الإنجاز" else "قابلية بدنية")
            frames.append(t[["phase", "week", "trial", "var", "value", "label", "unit", "kind"]])
    if not frames:
        return pd.DataFrame(columns=["phase", "week", "trial", "var", "value", "label", "unit", "kind", "order"])
    df = pd.concat(frames, ignore_index=True)
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    df = df.dropna(subset=["value"])
    df["order"] = [phase_order(p, w, total) for p, w in zip(df["phase"], df["week"])]
    return df


def single_subject_stats(trials, var, direction=0):
    """
    مقارنة محاولات القبلي بمحاولات البعدي للاعب واحد:
    الوسط والانحراف، الفرق، نسبة التحسن، اختبار t (Welch) ومان-ويتني بين المحاولات،
    حجم الأثر d، نسبة عدم التداخل PND، وأصغر تغير مهم SWC = 0.2 × انحراف القبلي.
    """
    from scipy import stats

    d = trials[trials["var"] == var]
    pre = d[d["phase"] == "قبلي"]["value"].to_numpy(float)
    post = d[d["phase"] == "بعدي"]["value"].to_numpy(float)
    r = {"n_pre": len(pre), "n_post": len(post)}
    r["m_pre"], r["sd_pre"] = (pre.mean(), pre.std(ddof=1) if len(pre) > 1 else np.nan) if len(pre) else (np.nan, np.nan)
    r["m_post"], r["sd_post"] = (post.mean(), post.std(ddof=1) if len(post) > 1 else np.nan) if len(post) else (np.nan, np.nan)
    r["diff"] = r["m_post"] - r["m_pre"]
    dirn = direction or (np.sign(r["diff"]) if not np.isnan(r["diff"]) else 0)
    # النسبة المئوية لا معنى لها إذا كانت القيم سالبة أو تعبر الصفر (مثل ميل الجذع)
    ok_pct = (len(pre) and len(post) and not np.isnan(r["m_pre"]) and np.all(pre > 0) and np.all(post > 0))
    r["pct"] = r["diff"] / r["m_pre"] * 100 * (direction or 1) if ok_pct else np.nan
    r["t"] = r["p_t"] = r["p_u"] = r["d"] = r["pnd"] = r["swc"] = np.nan
    if len(pre) >= 2 and len(post) >= 2 and (np.std(pre) > 0 or np.std(post) > 0):
        r["t"], r["p_t"] = stats.ttest_ind(post, pre, equal_var=False)
        sp = math.sqrt(((len(pre) - 1) * pre.var(ddof=1) + (len(post) - 1) * post.var(ddof=1)) / (len(pre) + len(post) - 2))
        r["d"] = r["diff"] / sp if sp > 0 else np.nan
    if len(pre) >= 3 and len(post) >= 3:
        r["p_u"] = stats.mannwhitneyu(post, pre, alternative="two-sided").pvalue
    if len(pre) and len(post) and dirn:
        best_pre = np.max(pre * dirn)
        r["pnd"] = float(np.mean(post * dirn > best_pre) * 100)
    if len(pre) >= 2:
        r["swc"] = 0.2 * r["sd_pre"]

    if np.isnan(r["diff"]):
        verdict = "لا يوجد قياس قبلي وبعدي"
    elif np.isnan(r["swc"]):
        verdict = "تغيّر (محاولة واحدة؛ لا يمكن فصله عن التذبذب)"
    elif abs(r["diff"]) <= r["swc"]:
        verdict = "ضمن التذبذب الطبيعي"
    elif direction == 0:
        verdict = "تغيّر واضح"
    elif r["diff"] * direction > 0:
        sig = (not np.isnan(r["p_t"]) and r["p_t"] < 0.05) or (not np.isnan(r["pnd"]) and r["pnd"] >= 90)
        verdict = "تحسن واضح" if sig else "تحسن محتمل"
    else:
        verdict = "تراجع"
    r["verdict"] = verdict
    return r


def single_subject_table(trials, bank):
    dirs = directions(bank)
    rows = []
    for (var, label, unit, kind), _ in trials.groupby(["var", "label", "unit", "kind"], sort=False):
        s = single_subject_stats(trials, var, dirs.get(var, 0))
        rows.append({"النوع": kind, "المتغير": label, "الوحدة": unit,
                     "محاولات قبلي": s["n_pre"], "قبلي: الوسط": s["m_pre"], "قبلي: الانحراف": s["sd_pre"],
                     "محاولات بعدي": s["n_post"], "بعدي: الوسط": s["m_post"], "بعدي: الانحراف": s["sd_post"],
                     "الفرق": s["diff"], "التحسن %": s["pct"], "حجم الأثر d": s["d"],
                     "p (t)": s["p_t"], "p (مان-ويتني)": s["p_u"], "PND %": s["pnd"], "الحكم": s["verdict"]})
    return pd.DataFrame(rows)


def sample_single(bank, seed=11):
    """لاعب واحد: 5 محاولات قبلية، متابعة في الأسابيع 2 و4 و6 (3 محاولات)، 5 محاولات بعدية."""
    rng = np.random.default_rng(seed)
    players = [["S1", "لاعب الحالة (تجريبي)", "تجريبية", "ذكر", 185, 76, 2006, "اليسرى"]]
    weeks = int(bank["settings"]["weeks"])
    plan = [("قبلي", "", 5)] + [("متابعة", w, 3) for w in (2, 4, 6)] + [("بعدي", "", 5)]
    mbase = {"approach_speed_mps": (6.2, 0.15, 0.35), "com_lowering_cm": (2.2, 0.8, 3.0),
             "knee_at_plant_deg": (151, 3, 9), "knee_min_deg": (133, 3, 7), "knee_at_takeoff_deg": (167, 3, 4),
             "contact_time_s": (0.235, 0.012, -0.035), "trunk_lean_plant_deg": (-3, 3, -8),
             "trunk_lean_takeoff_deg": (13, 3, -8), "free_knee_lift_pct": (-6, 2, 6),
             "free_knee_angle_deg": (98, 6, -10), "arms_raised_count": (0.6, 0.5, 1.0),
             "takeoff_angle_deg": (41, 2, 5), "hip_rise_cm": (80, 4, 9), "time_to_peak_s": (0.40, 0.02, 0.03),
             "hip_angle_at_peak_deg": (156, 4, 6), "knee_angle_at_peak_deg": (112, 6, -6)}
    tbase = {"T01": (47, 1.5, 5), "T02": (45, 1.5, 5), "T03": (240, 5, 12), "T04": (14.0, 0.3, 0.9),
             "T05": (3.38, 0.04, -0.12), "T06": (11.2, 0.4, 1.0), "T07": (1.40, 0.04, 0.18),
             "T08": (50, 2, -9), "T09": (9.0, 0.15, -0.4), "A01": (1.70, 0.02, 0.10)}
    tests, analyses = [], []
    for phase, week, n in plan:
        prog = 0.0 if phase == "قبلي" else 1.0 if phase == "بعدي" else week / (weeks + 1)
        for k in range(n):
            mv = {}
            for key, (m, sd, gain) in mbase.items():
                v = m + gain * prog + sd * rng.normal()
                if key == "arms_raised_count":
                    v = int(np.clip(round(v), 0, 2))
                mv[key] = round(float(v), 3 if key.endswith("_s") else 1)
            analyses.append(["S1", phase, week, k + 1] + list(mv.values()) +
                            [",".join(e["id"] for e in detect_errors(mv, "ذكر", bank))])
        for k in range(3 if phase != "متابعة" else 1):
            row = ["S1", phase, week]
            for t in bank["tests"]["رمز_الاختبار"]:
                m, sd, gain = tbase.get(t, (10, 0.3, 0.5))
                row.append(round(m + gain * prog + sd * rng.normal(), 2))
            tests.append(row)
    return players, tests, analyses


def _progress_png(trials, var, title):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    d = trials[trials["var"] == var]
    if d.empty:
        return None
    fig, ax = plt.subplots(figsize=(6.5, 2.8), dpi=150)
    ax.scatter(d["order"], d["value"], s=18, color="#9aa9bd", zorder=2, label="trials")
    m = d.groupby("order")["value"].mean()
    ax.plot(m.index, m.values, "-o", color="#1f4e78", lw=2, zorder=3, label="mean")
    ax.set_xlabel("week (0 = pre, last = post)")
    ax.set_title(title, fontsize=10)
    ax.grid(alpha=0.3)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def player_report_docx(player, table, trials, errors_by_phase, researcher=""):
    """تقرير دراسة الحالة للاعب واحد (Word، من اليمين لليسار)."""
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.6)
        b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1"); sec._sectPr.append(b)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    _para(doc, "تقرير تحليل اللاعب — الوثب العالي (فوسبري فلوب)", bold=True, size=15, align="center")
    _para(doc, f"اللاعب: {player.get('name', '')}  |  الرمز: {player.get('code', '')}  |  المجموعة: {player.get('group', '')}  |  "
               f"الطول: {player.get('height', '')} سم  |  الوزن: {player.get('weight', '')} كغم")
    if researcher:
        _para(doc, f"إعداد: {researcher}")

    def fmtv(v):
        try:
            f = float(v)
            return "—" if math.isnan(f) else f"{f:.2f}"
        except (TypeError, ValueError):
            return str(v)

    _para(doc, "أولاً: المقارنة القبلية/البعدية على مستوى المحاولات", bold=True, size=12)
    cols = ["المتغير", "الوحدة", "قبلي: الوسط", "بعدي: الوسط", "التحسن %", "حجم الأثر d", "PND %", "الحكم"]
    tbl = doc.add_table(rows=1, cols=len(cols))
    tbl.style = "Table Grid"
    _table_rtl(tbl)
    for i, c in enumerate(cols):
        r = tbl.rows[0].cells[i].paragraphs[0].add_run(c); r.bold = True
        r._r.get_or_add_rPr().append(OxmlElement("w:bCs"))
        _rtl_paragraph(tbl.rows[0].cells[i].paragraphs[0])
    widths = [Cm(4.6), Cm(1.4), Cm(1.8), Cm(1.8), Cm(1.7), Cm(1.7), Cm(1.5), Cm(2.6)]
    for i, wd in enumerate(widths):
        tbl.rows[0].cells[i].width = wd
    for _, row in table.iterrows():
        cells = tbl.add_row().cells
        for i, wd in enumerate(widths):
            cells[i].width = wd
        for i, c in enumerate(cols):
            cells[i].paragraphs[0].add_run(fmtv(row.get(c, "")) if c not in ("المتغير", "الوحدة", "الحكم") else str(row.get(c, "")))
            _rtl_paragraph(cells[i].paragraphs[0])
    _para(doc, "PND: نسبة محاولات البعدي الأفضل من أفضل محاولة قبلية. الحكم يقارن الفرق بأصغر تغير مهم (0.2 × انحراف القبلي).",
          size=8)

    _para(doc, "ثانياً: منحنيات التطور", bold=True, size=12)
    shown = 0
    for var in (["A01"] + [v for v in trials["var"].unique() if v != "A01"]):
        if shown >= 4 or var not in set(trials["var"]):
            continue
        lab = trials[trials["var"] == var]["label"].iloc[0]
        png = _progress_png(trials, var, str(var))
        if png:
            _para(doc, f"{lab} ({var})", bold=True)
            doc.add_picture(png, width=Inches(6.0))
            shown += 1

    _para(doc, "ثالثاً: الأخطاء الفنية عبر القياسات", bold=True, size=12)
    if errors_by_phase:
        for ph, errs in errors_by_phase.items():
            _para(doc, f"• {ph}: " + ("، ".join(f"{k} ({v})" for k, v in errs.items()) or "لا توجد أخطاء"))
    else:
        _para(doc, "لا توجد بيانات أخطاء مسجلة.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
